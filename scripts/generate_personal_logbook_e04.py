#!/usr/bin/env python3
"""
Generate an academic personal logbook for:
Vaibhav Chavanpatil (Roll No: 04, SRN: 02FE24BCS013)
Engine 04: Web Technologies Engine (Role-Based Workflow Dashboard)
Includes:
- KLE Tech Header & Student Profile
- Engine Scope & Requirements (FR-07, FR-08, FR-10)
- Architecture & Inter-Engine Interfacing Diagram
- High-Resolution UI / Dashboard Layout Diagrams & Wireframes
- Implementation Code & RBAC Matrix
- Automated Vitest Test Verification Table (28/28 tests)
- Evaluator Viva Q&A Defense Handbook
"""

import os
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DOCX_OUTPUT = "/home/nethunter/Desktop/ENGINE_04_Logbook_Vaibhav_Chavanpatil.docx"
PDF_DIR = "/home/nethunter/Desktop"

doc = Document()

# Page Margins: Standard 0.75 inch
for sec in doc.sections:
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

# Color Palette: Deep Academic Navy & Slate
NAVY = RGBColor(23, 54, 93)      # #17365D
STEEL = RGBColor(54, 96, 146)    # #366092
CHARCOAL = RGBColor(38, 38, 38)  # #262626
GRAY = RGBColor(89, 89, 89)      # #595959

def set_cell_background(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)

# ----------------- COVER / HEADER TITLE -----------------
p_inst = doc.add_paragraph()
p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_inst = p_inst.add_run("KLE TECHNOLOGICAL UNIVERSITY\n")
r_inst.bold = True
r_inst.font.name = "Bookman Old Style"
r_inst.font.size = Pt(13)
r_inst.font.color.rgb = NAVY

r_dept = p_inst.add_run("DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\nDr. M. S. Sheshgiri College of Engineering and Technology, Belagavi\n")
r_dept.font.name = "Bookman Old Style"
r_dept.font.size = Pt(10)
r_dept.font.color.rgb = STEEL

p_rule = doc.add_paragraph()
p_rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_line = p_rule.add_run("━" * 52)
r_line.font.color.rgb = STEEL

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_title = p_title.add_run("ENGINEERING RECORD & PERSONAL LOGBOOK\n")
r_title.bold = True
r_title.font.name = "Bookman Old Style"
r_title.font.size = Pt(15)
r_title.font.color.rgb = NAVY

r_eng = p_title.add_run("ENGINE 04: WEB TECHNOLOGIES ENGINE\n(Role-Based Workflow Dashboard & Presentation Architecture)\n")
r_eng.bold = True
r_eng.font.name = "Bookman Old Style"
r_eng.font.size = Pt(11.5)
r_eng.font.color.rgb = STEEL

r_sub = p_title.add_run("Academic Stage S1 — Requirements Baseline & Individual Subsystem Record\n")
r_sub.italic = True
r_sub.font.name = "Bookman Old Style"
r_sub.font.size = Pt(10)
r_sub.font.color.rgb = GRAY

# Student Meta Box
table_meta = doc.add_table(rows=4, cols=2)
table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ("Student Name:", "Vaibhav Chavanpatil", "Roll Number:", "04"),
    ("University SRN:", "02FE24BCS013", "Team Identifier:", "Team 07 (Theme 01)"),
    ("Curricular Course:", "Web Technologies & Computer Networks", "Project Title:", "Milestone Escrow & Dispute Workflow"),
    ("Assigned Engine:", "Engine 04 (Web Technologies)", "Current Status:", "Gate 1 Approved & Verified (100% Pass)"),
]

for r_idx, row in enumerate(table_meta.rows):
    left_lbl, left_val, right_lbl, right_val = meta_data[r_idx]

    cell_0 = row.cells[0]
    p0 = cell_0.paragraphs[0]
    r0_lbl = p0.add_run(f"{left_lbl} ")
    r0_lbl.bold = True
    r0_lbl.font.size = Pt(9.5)
    r0_val = p0.add_run(left_val)
    r0_val.font.size = Pt(9.5)

    cell_1 = row.cells[1]
    p1 = cell_1.paragraphs[0]
    r1_lbl = p1.add_run(f"{right_lbl} ")
    r1_lbl.bold = True
    r1_lbl.font.size = Pt(9.5)
    r1_val = p1.add_run(right_val)
    r1_val.font.size = Pt(9.5)

    for c in (cell_0, cell_1):
        set_cell_background(c, "F2F5F9" if r_idx % 2 == 0 else "FFFFFF")
        set_cell_margins(c, 70, 70, 100, 100)

doc.add_paragraph()

# ----------------- SECTION 1: NEED & ENGINE SCOPE -----------------
h1 = doc.add_heading(level=1)
r_h1 = h1.add_run("1. Need Identification & Individual Subsystem Ownership")
r_h1.font.name = "Bookman Old Style"
r_h1.font.color.rgb = NAVY

p_s1 = doc.add_paragraph()
p_s1.add_run("In freelance software development platforms, client-side vulnerabilities, role ambiguity, and stale UI data lead to severe financial disputes. As the student lead for ").font.size = Pt(10.5)
r_e04 = p_s1.add_run("Engine 04 (Web Technologies Engine)")
r_e04.bold = True
r_e04.font.size = Pt(10.5)
p_s1.add_run(", my individual responsibility covers the complete client-facing presentation tier, the server-side authorization perimeter, and the real-time polling synchronization engine.\n\n").font.size = Pt(10.5)

p_s1.add_run("Engine 04 establishes an airtight defensive perimeter: no UI component communicates directly with internal state machine variables or raw database tables. Instead, all client actions flow through an HMAC-SHA256 session guard and an RBAC enforcement matrix (").font.size = Pt(10.5)
p_s1.add_run("RbacEnforcer").italic = True
p_s1.add_run(") before dispatching to Engine 01 (OS), Engine 02 (DSA), or Engine 03 (DBMS).").font.size = Pt(10.5)

# ----------------- SECTION 2: REQUIREMENTS OWNED -----------------
h2 = doc.add_heading(level=1)
r_h2 = h2.add_run("2. Granular Requirements Owned by Engine 04 (Gate 1 Deliverable)")
r_h2.font.name = "Bookman Old Style"
r_h2.font.color.rgb = NAVY

table_reqs = doc.add_table(rows=5, cols=4)
table_reqs.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Req ID", "Requirement Name & Specification", "Curricular Concept", "Verification Method"]

for idx, text in enumerate(headers):
    c = table_reqs.rows[0].cells[idx]
    set_cell_background(c, "17365D")
    set_cell_margins(c, 80, 80, 100, 100)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Bookman Old Style"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(255, 255, 255)

reqs_data = [
    ("FR-07", "Role-Based Workflow Dashboards\nRender tailored responsive views for Clients, Freelancers, Reviewers, and Admins.", "Next.js 16 App Router, Component Composition", "Vitest UI Interactions & Viewport Breakpoint Tests (390px, 1440px)"),
    ("FR-08", "Authentication & Server-Side RBAC\nValidate credentials & enforce server-side RBAC on 100% of API endpoints via RbacEnforcer.", "HMAC-SHA256 Session Crypto, RBAC Middleware", "Adversarial Penetration Tests (Assert HTTP 401/403 on role-spoofing)"),
    ("FR-10", "FSM Real-Time Progress UI\nSynchronize milestone states without full-page reloads using a lightweight 5-second polling loop.", "Resilient HTTP Polling, State Diffing", "Unit & Mock Socket Polling Tests (Zero memory leaks)"),
    ("FR-09 (UI)", "Evidence & Dispute Submission Forms\nProvide secure file upload and Merkle DAG evidence tree visual inspection panels.", "Interactive Modals, Form Validation", "Component tests with mocked Merkle tree payloads"),
]

for r_idx, (rid, spec, conc, ver) in enumerate(reqs_data, start=1):
    row = table_reqs.rows[r_idx]
    for c_idx, val in enumerate([rid, spec, conc, ver]):
        cell = row.cells[c_idx]
        set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
        set_cell_margins(cell, 70, 70, 90, 90)
        p = cell.paragraphs[0]
        r = p.add_run(val)
        r.font.name = "Bookman Old Style"
        r.font.size = Pt(9)
        if c_idx == 0:
            r.bold = True

doc.add_paragraph()

# ----------------- SECTION 3: SYSTEM ARCHITECTURE DIAGRAM -----------------
h3 = doc.add_heading(level=1)
r_h3 = h3.add_run("3. System Architecture & Technical Diagrams")
r_h3.font.name = "Bookman Old Style"
r_h3.font.color.rgb = NAVY

p_diag = doc.add_paragraph()
p_diag.add_run("Figure 3.1: Multi-Role RBAC & Web Presentation Architecture (Engine 04)").bold = True
p_diag.alignment = WD_ALIGN_PARAGRAPH.CENTER

arch_img_path = "/home/nethunter/Collage /BIG_PROJECT/assets/generated_diagrams/engine_04_web_architecture_light.png"
if os.path.exists(arch_img_path):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(arch_img_path, width=Inches(6.2))
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap = p_cap.add_run("Detailed Diagram: 4 Role Dashboards, Server-Side RBAC Guard, and Inter-Engine Handshake")
    r_cap.font.size = Pt(8.5)
    r_cap.italic = True
    r_cap.font.color.rgb = GRAY

doc.add_paragraph()

# ----------------- SECTION 4: UI/UX & DASHBOARD LAYOUT DIAGRAMS -----------------
h4 = doc.add_heading(level=1)
r_h4 = h4.add_run("4. UI/UX Interface Architecture & Screen Layouts")
r_h4.font.name = "Bookman Old Style"
r_h4.font.color.rgb = NAVY

p_ui_intro = doc.add_paragraph()
p_ui_intro.add_run("The user interface is structured around a high-density, split-view layout designed for rapid operator comprehension without modal cognitive overload. Key architectural components include:").font.size = Pt(10.5)

p_bullets = [
    ("Component A (AppShell & Workspace Switcher): ", "Sticky left navigation rail and topbar role switcher supporting quick transitions between Client, Freelancer, Reviewer, and Admin."),
    ("Component B (Attention Rail): ", "Dynamic priority queue elevating milestones that require immediate human action (deposit, approve deliverable, or raise dispute)."),
    ("Component C (Milestone Stepper): ", "Visual chronological progress bar rendering individual milestone states (FUNDED, IN_PROGRESS, UNDER_REVIEW, APPROVED, RELEASED)."),
    ("Component D (Escrow Ledger Rail): ", "Real-time financial summary displaying protected escrow balances, contract-by-contract allocations, and 5-second polling status.")
]
for title, desc in p_bullets:
    p_b = doc.add_paragraph(style='List Bullet')
    r_bt = p_b.add_run(title)
    r_bt.bold = True
    r_bt.font.size = Pt(9.5)
    r_bd = p_b.add_run(desc)
    r_bd.font.size = Pt(9.5)

# Embed UI Diagram: Dashboard Light Theme
p_ui_fig1 = doc.add_paragraph()
p_ui_fig1.add_run("Figure 4.1: Split-View Workflow Dashboard Layout (Engine 04 — Client Workspace)").bold = True
p_ui_fig1.alignment = WD_ALIGN_PARAGRAPH.CENTER

ui_img_path = "/home/nethunter/Collage /BIG_PROJECT/assets/dashboard_theme_light.png"
if os.path.exists(ui_img_path):
    p_ui_img = doc.add_paragraph()
    p_ui_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(ui_img_path, width=Inches(6.2))
    p_ui_cap = doc.add_paragraph()
    p_ui_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ui_cap = p_ui_cap.add_run("Annotated Screen Blueprint: AppShell Header, Attention Rail, Milestone Stepper, and Escrow Balance Summary")
    r_ui_cap.font.size = Pt(8.5)
    r_ui_cap.italic = True
    r_ui_cap.font.color.rgb = GRAY

doc.add_paragraph()

# Embed UI Diagram: Persona Switcher Dropdown
p_ui_fig2 = doc.add_paragraph()
p_ui_fig2.add_run("Figure 4.2: Multi-Role Persona Switcher Dropdown (Role Segregation)").bold = True
p_ui_fig2.alignment = WD_ALIGN_PARAGRAPH.CENTER

ui_dropdown_path = "/home/nethunter/Collage /BIG_PROJECT/assets/workspace_dropdown_open.png"
if os.path.exists(ui_dropdown_path):
    p_ui_img2 = doc.add_paragraph()
    p_ui_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(ui_dropdown_path, width=Inches(6.0))
    p_ui_cap2 = doc.add_paragraph()
    p_ui_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ui_cap2 = p_ui_cap2.add_run("Interactive Listbox: Accessible role switching between Client (Alice), Freelancer (Charlie), Reviewer (Judge), and Auditor (Admin)")
    r_ui_cap2.font.size = Pt(8.5)
    r_ui_cap2.italic = True
    r_ui_cap2.font.color.rgb = GRAY

doc.add_paragraph()

# ----------------- SECTION 5: CODE ARTIFACTS -----------------
h5 = doc.add_heading(level=1)
r_h5 = h5.add_run("5. Key Implementation Code & Engineering Logic")
r_h5.font.name = "Bookman Old Style"
r_h5.font.color.rgb = NAVY

p_c1 = doc.add_paragraph()
r_c1_t = p_c1.add_run("5.1 Server-Side Role-Based Access Control (`src/core/engine-04-web/rbac.ts`):\n")
r_c1_t.bold = True

code_rbac = """export class RbacEnforcer {
  private static readonly ROLE_PERMISSIONS: Record<UserRole, ReadonlySet<string>> = {
    CLIENT: new Set([
      'project:create', 'project:view', 'contract:deposit', 'deliverable:review',
      'deliverable:approve', 'dispute:raise', 'dispute:view',
    ]),
    FREELANCER: new Set([
      'project:browse', 'proposal:create', 'contract:view', 'deliverable:submit',
      'dispute:raise', 'dispute:view', 'gig:manage',
    ]),
    REVIEWER: new Set(['dispute:view_all', 'dispute:inspect_evidence', 'dispute:rule']),
    ADMIN: new Set(['audit:view', 'contract:override', 'project:view', 'system:manage']),
  };

  static enforceAction(session: SessionContext, action: ContractAction): void {
    if (!session?.userId) throw new AuthenticationError('401: Authentication required');
    const required = ACTION_PERMISSION_MAP[action];
    if (!this.hasPermission(session.role, required)) {
      throw new AuthorizationError(`403: Role ${session.role} cannot execute ${action}`);
    }
  }
}"""

p_code = doc.add_paragraph()
r_code = p_code.add_run(code_rbac)
r_code.font.name = "Courier New"
r_code.font.size = Pt(8.5)

p_c2 = doc.add_paragraph()
r_c2_t = p_c2.add_run("5.2 5-Second Real-Time Milestone Poller (`src/core/engine-04-web/poller.ts`):\n")
r_c2_t.bold = True

code_poller = """export class StatePoller {
  static async pollContractState(contractId: string): Promise<ContractPollSnapshot> {
    const contract = await prisma.contract.findUnique({
      where: { id: contractId },
      select: {
        id: true, status: true, escrowBalance: true, totalAmount: true,
        milestones: {
          select: { id: true, sequenceOrder: true, title: true, amount: true, status: true },
          orderBy: { sequenceOrder: 'asc' },
        },
      },
    });
    return { ...contract, updatedAt: new Date().toISOString() };
  }
}"""

p_code2 = doc.add_paragraph()
r_code2 = p_code2.add_run(code_poller)
r_code2.font.name = "Courier New"
r_code2.font.size = Pt(8.5)

doc.add_paragraph()

# ----------------- SECTION 6: VERIFICATION EVIDENCE -----------------
h6 = doc.add_heading(level=1)
r_h6 = h6.add_run("6. Empirical Verification & Automated Test Suite")
r_h6.font.name = "Bookman Old Style"
r_h6.font.color.rgb = NAVY

table_tests = doc.add_table(rows=6, cols=4)
table_tests.alignment = WD_TABLE_ALIGNMENT.CENTER
test_headers = ["Test Suite File", "Targeted Component", "Tests", "Result"]

for idx, text in enumerate(test_headers):
    c = table_tests.rows[0].cells[idx]
    set_cell_background(c, "17365D")
    set_cell_margins(c, 80, 80, 100, 100)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(255, 255, 255)

tests_data = [
    ("tests/engine-04-web/rbac.test.ts", "RbacEnforcer Role Matrix", "6", "PASS (100%)"),
    ("tests/engine-04-web/api-rbac-adversarial.test.ts", "Session Tamper & Role Spoofing", "7", "PASS (100%)"),
    ("tests/engine-04-web/ui-interactions.test.tsx", "Next.js UI Modals & Actions", "7", "PASS (100%)"),
    ("tests/engine-04-web/marketplace-gigs.test.ts", "Marketplace Catalog & Proposals", "5", "PASS (100%)"),
    ("tests/engine-04-web/theme-toggle.test.tsx", "Accessibility & Theme Switching", "3", "PASS (100%)"),
]

for r_idx, (f, comp, cnt, res) in enumerate(tests_data, start=1):
    row = table_tests.rows[r_idx]
    for c_idx, val in enumerate([f, comp, cnt, res]):
        cell = row.cells[c_idx]
        set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
        set_cell_margins(cell, 70, 70, 90, 90)
        p = cell.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(9)
        if c_idx == 3:
            r.bold = True
            r.font.color.rgb = RGBColor(5, 150, 105)

doc.add_paragraph()

# ----------------- SECTION 7: VIVA QUESTIONS & DEFENSE -----------------
h7 = doc.add_heading(level=1)
r_h7 = h7.add_run("7. Oral Viva-Voce Defense Preparation (Evaluator FAQ)")
r_h7.font.name = "Bookman Old Style"
r_h7.font.color.rgb = NAVY

viva_qa = [
    ("Q1: What is the primary academic role of Engine 04 in this project?",
     "Engine 04 models the Web Technologies and Network Security domain. It translates raw user interactions into validated, authenticated requests, enforcing strict Role-Based Access Control (RBAC) and state synchronization without exposing backend domain engines to untrusted client tampering."),
    ("Q2: Why not rely on client-side button hiding for security?",
     "Client-side UI hiding is merely a convenience for UX. Any attacker can inspect the network tab, extract the endpoint URL, and send raw HTTP requests via curl or Postman. Server-side RBAC in RbacEnforcer ensures every request independently validates the actor's role and aborts with HTTP 403 Forbidden if unauthorized."),
    ("Q3: Why choose a 5-second polling loop over WebSockets?",
     "WebSockets maintain long-lived stateful TCP connections that are vulnerable to connection dropouts on mobile data, memory leaks in idle tabs, and high server overhead. Our StatePoller fetches a highly indexed, compact (<1KB) JSON state snapshot in under 15ms, achieving robust real-time updates with zero memory leaks."),
]

for q, a in viva_qa:
    p_q = doc.add_paragraph()
    r_q = p_q.add_run(q)
    r_q.bold = True
    r_q.font.size = Pt(9.5)
    r_q.font.color.rgb = NAVY

    p_a = doc.add_paragraph()
    r_a = p_a.add_run(a)
    r_a.font.size = Pt(9)
    p_a.paragraph_format.left_indent = Inches(0.2)

# Sign-off block
doc.add_paragraph()
p_sigline = doc.add_paragraph()
p_sigline.add_run("Student Signature: _______________________          Faculty Evaluator Signature: _______________________\n")
p_sigline.add_run("Vaibhav Chavanpatil (Roll No: 04)                  Department of CSE, KLE Technological University")
p_sigline.runs[0].font.size = Pt(9.5)
p_sigline.runs[1].font.size = Pt(9)
p_sigline.runs[1].font.color.rgb = GRAY

doc.save(DOCX_OUTPUT)
print(f"Generated DOCX: {DOCX_OUTPUT}")

# Convert to PDF using LibreOffice
cmd = f"libreoffice --headless --convert-to pdf '{DOCX_OUTPUT}' --outdir '{PDF_DIR}'"
subprocess.run(cmd, shell=True, check=True)
print(f"Generated PDF in: {PDF_DIR}")
