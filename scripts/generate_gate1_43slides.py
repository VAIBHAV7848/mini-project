import sys
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    NAVY = RGBColor(30, 58, 138)       # #1E3A8A
    DARK_BLUE = RGBColor(26, 54, 93)   # #1A365D
    SLATE = RGBColor(100, 116, 139)    # #64748B
    TEXT_DARK = RGBColor(30, 41, 59)   # #1E293B
    BG_LIGHT = RGBColor(248, 250, 252) # #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)  # White
    BORDER_CLR = RGBColor(226, 232, 240) # #E2E8F0
    ACCENT_TEAL = RGBColor(13, 148, 136) # #0D9488
    ACCENT_BLUE = RGBColor(37, 99, 235)  # #2563EB
    ACCENT_PURPLE = RGBColor(109, 40, 217) # #6D28D9
    ACCENT_AMBER = RGBColor(217, 119, 6)   # #D97706
    ACCENT_RED = RGBColor(225, 29, 72)     # #E11D48

    def add_blank_slide():
        slide = prs.slides.add_slide(blank_layout)
        # Add background rectangle
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_LIGHT
        bg.line.color.rgb = BG_LIGHT
        return slide

    def add_header(slide, title, subtitle=None, badge=None):
        top_y = 0.35
        if badge:
            # Add circular or pill badge for deliverable number
            badge_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(top_y), Inches(0.35), Inches(0.35))
            badge_box.fill.solid()
            badge_box.fill.fore_color.rgb = NAVY
            badge_box.line.color.rgb = NAVY
            tf = badge_box.text_frame
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.text = str(badge)
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
            title_left = 1.05
        else:
            title_left = 0.6

        # Title
        tb = slide.shapes.add_textbox(Inches(title_left), Inches(top_y - 0.05), Inches(11.5), Inches(0.45))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = DARK_BLUE

        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.size = Pt(11)
            p2.font.color.rgb = SLATE

    def add_footer(slide, text="Milestone-Based Escrow & Evidence-Based Dispute Resolution Workflow | Gate 1"):
        tb = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12.133), Inches(0.3))
        tf = tb.text_frame
        tf.margin_top = Inches(0)
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(8.5)
        p.font.color.rgb = SLATE

    def add_card(slide, left, top, width, height, bg_clr=CARD_BG, border_clr=BORDER_CLR):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_clr
        card.line.color.rgb = border_clr
        card.line.width = Pt(1)
        return card

    def add_table(slide, left, top, width, height, headers, rows, col_widths=None):
        table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(left), Inches(top), Inches(width), Inches(height))
        tbl = table_shape.table

        # set column widths if provided
        if col_widths:
            for i, w in enumerate(col_widths):
                tbl.columns[i].width = Inches(w)

        # Header Row
        for i, h in enumerate(headers):
            cell = tbl.cell(0, i)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY
            tf = cell.text_frame
            tf.word_wrap = True
            tf.margin_top = Inches(0.06)
            tf.margin_bottom = Inches(0.06)
            tf.margin_left = Inches(0.08)
            tf.margin_right = Inches(0.08)
            p = tf.paragraphs[0]
            p.text = h
            p.font.bold = True
            p.font.size = Pt(10.5)
            p.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for r_idx, row in enumerate(rows):
            is_alt = (r_idx % 2 == 1)
            row_bg = RGBColor(241, 245, 249) if is_alt else RGBColor(255, 255, 255)
            for c_idx, val in enumerate(row):
                cell = tbl.cell(r_idx + 1, c_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = row_bg
                tf = cell.text_frame
                tf.word_wrap = True
                tf.margin_top = Inches(0.05)
                tf.margin_bottom = Inches(0.05)
                tf.margin_left = Inches(0.08)
                tf.margin_right = Inches(0.08)
                p = tf.paragraphs[0]
                p.text = str(val)
                p.font.size = Pt(9.5)
                p.font.color.rgb = TEXT_DARK
        return table_shape

    print("Setup helper functions complete.")
    return prs

prs = create_presentation()
print("Success initializing test.")
